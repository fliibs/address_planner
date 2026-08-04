"""Generate a compact, deterministic DOCX view of an address-map JSON tree."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


def _value(node: dict, key: str, fallback: str = "-") -> str:
    value = node.get(key, fallback)
    return fallback if value is None or value == "" else str(value)


def _add_node(document: Document, node: dict, numbering: tuple[int, ...]) -> None:
    label = ".".join(str(part) for part in numbering)
    name = _value(node, "name", "unnamed")
    heading_level = min(len(numbering), 3)
    document.add_heading(f"{label} {name}", level=heading_level)

    details = document.add_table(rows=0, cols=2)
    details.style = "Table Grid"
    for title, key in (
        ("Type", "type"),
        ("Start address", "start_addr"),
        ("End address", "end_addr"),
        ("Size", "size"),
        ("Description", "description"),
    ):
        cells = details.add_row().cells
        cells[0].text = title
        cells[1].text = _value(node, key)

    fields = node.get("fields") or []
    if fields:
        document.add_paragraph("Fields")
        table = document.add_table(rows=1, cols=8)
        table.style = "Table Grid"
        headers = (
            "Name",
            "Position",
            "Size",
            "Software access",
            "Hardware access",
            "External",
            "Default value",
            "Description",
        )
        for cell, title in zip(table.rows[0].cells, headers):
            cell.text = title
        for field in fields:
            cells = table.add_row().cells
            values = (
                _value(field, "name"),
                _value(field, "Position"),
                _value(field, "size"),
                _value(field, "Software Access"),
                _value(field, "Hardware Access"),
                _value(field, "External"),
                _value(field, "Default Value"),
                _value(field, "description"),
            )
            for cell, value in zip(cells, values):
                cell.text = value

    for index, child in enumerate(node.get("children") or [], start=1):
        _add_node(document, child, (*numbering, index))


def build_address_map_document(root: dict, output_path) -> Path:
    """Write one DOCX report for the supplied address-map JSON root."""

    if not isinstance(root, dict):
        raise TypeError("address-map document root must be a dictionary")

    destination = Path(output_path).expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading("Address Map", level=0)
    _add_node(document, root, (1,))
    document.save(destination)
    return destination


__all__ = ["build_address_map_document"]
