from copy               import deepcopy
from functools          import reduce
from tkinter            import Tcl
from .AddressLogicRoot  import *
from .GlobalValues      import *
from .ralf_parser.ralf_parse import build_addrspace,py_dict
from .address_planner_rtl.MatrixCFG import *
from .gen_doc.doc import *

import os
import builtins
import json
import shutil
import re
from pathlib import Path
import openpyxl

class AddressSpace(AddressLogicRoot):

    def __init__(self,name,size=None,description='',path='./'):
        super().__init__(name=name,description=description,path=path)
        self.size           = size
        self.sub_space_list = []
        self.offset         = 0
        self._next_offset   = 0
        self.matrix_list    = []
        #self.module_name    = name
        #self.module_name      = ''
        #self.name           = name
        #self.start_address  = start_address
        #self.end_address    = self.start_address + self.size - 1
        #self.description    = description
        #self.path           = path
        #self.father         = None



    @property
    def bit_offset(self):
        return self.offset*8
    
    @property
    def bit_size(self):
        return self.size*8

    @property
    def global_size(self):
        return self.size*8 if self.father is None else self.father.bit_size

    @property
    def global_offset(self):
        return 0 if self.father is None else self.father.global_offset + self.bit_offset

    @property
    def global_start_address(self):
        return self.global_offset

    @property
    def global_end_address(self):
        return self.global_offset + self.bit_size - 1

    @property
    def start_address(self):
        return self.bit_offset

    @property
    def end_address(self):
        return self.bit_offset + self.bit_size - 1

    @property
    def hex_offset(self):
        hex_value = hex(self.bit_offset)
        if hex_value == '0x0':
            return '\'h0'
        else:
            return '\'h'+hex_value.lstrip('0x')
    
    @property
    def sorted_subspace_list(self):
        return sorted(self.sub_space_list, key=lambda x: x.bit_offset)
    
    @property
    def filled_sub_space_list(self):
        from .RegSpace import RegSpace
        reserve_idx = 0
        res = []
        previous_space = None

        if self.sorted_subspace_list == []:
            return []

        if self.sorted_subspace_list[0].bit_offset != 0:
            filled_subspace = RegSpace(name = f'Reserved_{reserve_idx}',size = self.sorted_subspace_list[0].offset)
            filled_subspace.offset = 0
            filled_subspace.father = self
            res.append(filled_subspace)
            reserve_idx += 1

        for space in self.sorted_subspace_list:
            if previous_space != None and space.start_address > previous_space.end_address + 1:
                filled_subspace = RegSpace(name = f'Reserved_{reserve_idx}',size = int((space.start_address - previous_space.end_address - 1)/8))
                filled_subspace.offset = int(previous_space.end_address/8) + 1
                filled_subspace.father = self
                res.append(filled_subspace)
                reserve_idx += 1
            res.append(space)
            previous_space = space

        if self.sorted_subspace_list[-1].end_address < self.bit_size - 1:
            filled_subspace = RegSpace(name = f'Reserved_{reserve_idx}',size = int((self.bit_size - previous_space.end_address - 1)/8))
            filled_subspace.offset = int(previous_space.end_address/8) + 1
            filled_subspace.father = self
            res.append(filled_subspace)
        
        return res
        

    def add(self,sub_space,offset,name=None):
        sub_space_copy = deepcopy(sub_space)
        sub_space_copy.offset = offset
        sub_space_copy.father = self
        # sub_space_copy.module_name = name
        sub_space_copy.module_name = sub_space_copy.module_name if name==None else name
        if not Options.MultiPortOption:
            if not self.inclusion_detect(sub_space_copy):
                raise ValueError('Sub space %s is not included in address space %s' %(sub_space_copy.module_name,self.module_name))

            for exist_space in self.sub_space_list:
                if self.collision_detect(exist_space,sub_space_copy):
                    raise ValueError('Address overlap: sub space %s(%s to %s) and current sub space %s(%s to %s) conflict.' \
                        % (sub_space_copy.module_name,hex(sub_space_copy.start_address),hex(sub_space_copy.end_address),exist_space.module_name,hex(exist_space.start_address),hex(exist_space.end_address)))
        self.sub_space_list.append(sub_space_copy)
        self._next_offset = offset + sub_space.size


    def add_incr(self,sub_space,name):
        self.add(sub_space=sub_space,offset=self._next_offset,name=name)


    def add_ralf(self,ralf_file,offset,name=None):
        ralf_path = Path(ralf_file).expanduser().resolve()
        if not ralf_path.is_file():
            raise FileNotFoundError(f'RALF input does not exist: {ralf_path}')
        with ralf_path.open('r', encoding='utf-8') as f:
            env_tcl_code = f.read()

        env_tcl_code = env_tcl_code.replace("[","").replace("]","")
        env_tcl_code = re.sub(r'\([^)]*\)', '', env_tcl_code)

        tcl_interpreter = Tcl()
        parser_tcl = Path(__file__).resolve().parent / 'ralf_parser' / 'ralf_parser.tcl'
        tcl_interpreter.call('source', str(parser_tcl))
        tcl_interpreter.eval(env_tcl_code)
        # A RALF description can intentionally contain alternate register views at
        # one address.  Keep that source-level information during import, while
        # restoring the normal overlap policy before returning to the caller.
        original_multiport = Options.MultiPortOption
        Options.MultiPortOption = True
        try:
            reg_copy = build_addrspace(tcl_interpreter)
        finally:
            Options.MultiPortOption = original_multiport
        self.add(reg_copy, offset, name)
        return reg_copy

    def add_matrix(self, matrix, name=None, attr=None):
        matrix_copy = deepcopy(matrix)
        matrix_copy.father = self
        matrix_copy.module_name = matrix_copy.module_name if name is None else name
        matrix_copy.attr = matrix_copy.attr if attr is None else attr
        self.matrix_list.append(matrix_copy)

    def update_matrix(self, sub_space, name):
        for matrix in self.matrix_list:
            matrix.update(sub_space, name)

    def hex_transform(self, value):
        if not isinstance(value, str):
            hex_value = hex(value)
        if hex_value == '0x0':
            return '\'h0'
        else:
            return '\'h'+hex_value.lstrip('0x')

    # def reg_bit_detect(self, sub_space):
    #     if sub_space.bit%32!=0: return False
    #     else:                   return True

    def collision_detect(self,space_A,space_B):
        if      (space_A.start_address <= space_B.start_address ) and (space_B.start_address <= space_A.end_address ): return True
        elif    (space_A.start_address <= space_B.end_address   ) and (space_B.end_address   <= space_A.end_address ): return True
        elif    (space_B.start_address <= space_A.start_address ) and (space_A.start_address <= space_B.end_address ): return True
        elif    (space_B.start_address <= space_A.end_address   ) and (space_A.end_address   <= space_B.end_address ): return True
        else:                                                                                                return False

    def inclusion_detect(self,other):
        return True if (self.start_address <= other.start_address) and (other.end_address <= self.end_address) else False
    
    def intr_detect(self,space):
        if space.reg_type==Intr and space.bit!=IntrBitWidth.IntrFull.value:             return False 
        elif space.reg_type==IntrMask and space.bit!=IntrBitWidth.IntrFull.value:   return False
        else:                                                                       return True

    def search_field(self, reg_name, field_name):
        for sub_space in self.sub_space_list:
            for field in sub_space.field_list:
                if reg_name == sub_space.module_name and field_name == field.name:
                    self.lock_inclusion_detect(field)
                    return [sub_space, field]
        
        raise Exception(f'lock field {reg_name}_{field_name} not exist')
    
    def search_magic(self, reg_name):
        for sub_space in self.sub_space_list:
            if reg_name == sub_space.module_name:
                self.magic_inclusion_detect(sub_space)
                return sub_space
        raise Exception(f'magic register {reg_name} not exist')

    
    def check_list(self, other):
        if not isinstance(other, list):     raise Exception("input must be a List type")

    def lock_inclusion_detect(self, field):
        from .Field import LockField
        if not isinstance(field, LockField):
            raise Exception(f'field in lock list is not LockField Type')
        
    def magic_inclusion_detect(self, reg):
        from .Field import MagicNumber
        if not isinstance(reg.field_list[0], MagicNumber):
            raise Exception(f'field in magic list is not MagicNumber Type')
        

    #########################################################################################
    # output generate
    #########################################################################################
             
    # def report_html(self):
    #     text = self.report_from_template(APG_HTML_FILE_ADDR_SPACE)
    #     os.makedirs(os.path.dirname(self.html_path), exist_ok=True)
    #     with open(self.html_path,'w') as f:
    #         f.write(text)
    #     #for ss in self.sub_space_list:
    #     #    ss.report_html()


    def report_chead(self):
        chead_name_list = self.report_chead_core()
        chead_name_list += [self.chead_global_name]
        self.report_chead_global_core()
        # Preserve recursive/model order while removing duplicate includes.  A
        # set-to-list conversion is randomized across Python processes.
        chead_name_list = list(dict.fromkeys(chead_name_list))
        with open(os.path.join(self._chead_dir,'all.h'),'w') as f:
            for chead_name in chead_name_list:
                f.write("#include \"%s\"\n" % chead_name)

    def report_chead_core(self):
        if self.sub_space_list == []:
            return []
        else:
            chead_name_list = [self.chead_name]
            text = self.report_from_template(APG_CHEAD_FILE_ADDR_SPACE,{'head_type':'c'})
            os.makedirs(os.path.dirname(self.chead_path), exist_ok=True)
            with open(self.chead_path,'w') as f:
                f.write(text)
            for ss in self.sub_space_list:
                chead_name_list += ss.report_chead_core()
            return chead_name_list
        
    def report_chead_global_core(self, file=None):
        if self.father == None: 
            os.makedirs(os.path.dirname(self.chead_global_path), exist_ok=True)
            file = open(self.chead_global_path,'w')
        if self.sub_space_list == []:
            text = self.report_from_template(APG_CHEAD_GLB_FILE_REG_SPACE)
            file.write(text)
        else:
            for ss in self.filled_sub_space_list:
                ss.report_chead_global_core(file)


    # report v head.==============================================
    def report_vhead(self):
        vhead_name_list = self.report_vhead_core()
        vhead_name_list += [self.vhead_global_name]
        self.report_vhead_global_core()
        vhead_name_list = list(dict.fromkeys(vhead_name_list))
        with open(os.path.join(self._vhead_dir,'all.vh'),'w') as f:
            for vhead_name in vhead_name_list:
                f.write("`include \"%s\"\n" % vhead_name)

        
    def report_vhead_core(self):
        if self.sub_space_list == []:
            return []
        else:
            vhead_name_list = [self.vhead_name]
            text = self.report_from_template(APG_VHEAD_FILE_ADDR_SPACE,{'head_type':'v'})
            os.makedirs(os.path.dirname(self.vhead_global_path), exist_ok=True)
            with open(self.vhead_path,'w') as f:
                f.write(text)
            for ss in self.sub_space_list:
                vhead_name_list += ss.report_vhead_core()
            return vhead_name_list
    
    def report_vhead_global_core(self, file=None):
        if self.father == None: 
            os.makedirs(os.path.dirname(self.vhead_path), exist_ok=True)
            file = open(self.vhead_global_path,'w')
        if self.sub_space_list == []:
            text = self.report_from_template(APG_VHEAD_GLB_FILE_REG_SPACE)
            file.write(text)
        else:
            for ss in self.filled_sub_space_list:
                ss.report_vhead_global_core(file)
        

    # report and check ralf ==============================================
    def report_ralf(self):
        output_path = self._ralf_dir+'/'
        self.recursive_report_ralf_core(output_path)

    def recursive_report_ralf_core(self, output_dir):
        for ss in self.sub_space_list:
            if hasattr(ss,'report_ralf_core'):
                ss.report_ralf_core(output_dir)
            else:
                ss.recursive_report_ralf_core(output_dir)


    # report and check json ==========================================
    def report_json(self, gen_doc=False):
        json_list= [self.report_json_core()]
        jtext = json.dumps(json_list, ensure_ascii=False, indent=2)
        if not os.path.exists(self._html_dir):  os.makedirs(self._html_dir) 
        with open(self.json_path, 'w') as f:
            f.write(jtext)

        root_path = os.environ.get('PARSER_PATH')
        if root_path == None:
            print("[Warning] no env PARSER_PATH, will not generate index.html.")
        else:
            index_path      = "apv_html/apv_html_main/py/index.html"
            script_path     = "apv_html/apv_html_main/py/main.py"
            full_index_path = os.path.join(root_path, index_path)
            full_script_path = os.path.join(root_path, script_path)
            shutil.copy(full_index_path, self._html_dir)
            print(f'python3 {full_script_path} -html {self.index_html_path} -json {self.json_path}')
            os.system(f'python3 {full_script_path} -html {self.index_html_path} -json {self.json_path}')

        if gen_doc:
            print("generate doc ! ")
            # Create a new document
            doc = Document()

            # Style settings
            style = doc.styles['Normal']
            font = style.font
            font.name = 'Times New Roman'
            font.size = Pt(12)

            heading1_style = doc.styles['Heading 1']
            heading1_font = heading1_style.font
            heading1_font.name = 'Times New Roman'
            heading1_font.size = Pt(20)

            heading2_style = doc.styles['Heading 2']
            heading2_font = heading2_style.font
            heading2_font.name = 'Times New Roman'
            heading2_font.size = Pt(18)

            heading3_style = doc.styles['Heading 3']
            heading3_font = heading3_style.font
            heading3_font.name = 'Times New Roman'
            heading3_font.size = Pt(16)

            # Add the main title to the document
            title = "Address map"
            doc.add_heading(title, level=0)
            doc.add_paragraph("")

            # Parse the JSON data
            doc_path = os.path.join(self._html_dir, "doc.docx")
            data = []
            with open(self.json_path, 'r') as file:
                data = json.load(file)

            # Traverse the JSON data
            traverse_json_with_numbering(doc, data[0])

            # Save the document
            doc.save(doc_path)
        

    def report_json_core(self):
        json_dict={}
        json_dict["key"]        = ADD_KEY()
        json_dict["type"]       = "sys"
        json_dict["name"]       = self.module_name
        # if self.father == None:
        json_dict["start_addr"] = hex(int(self.global_start_address/8))
        json_dict["end_addr"]   = hex(int(self.global_end_address/8))
        
        json_dict["size"]       = ConvertSize(self.size, is_byte=True)
        json_dict["description"]= self.description
        json_dict["children"]   = [c.report_json_core() for c in self.sorted_subspace_list]
        return json_dict

    def report_sqlite(self, database_path=None, *, schema_version=None):
        """Write the normalized SQLite report used by the single-HTML viewer.

        Schema v2 is the default; pass ``schema_version=SCHEMA_VERSION_V1``
        for the legacy layout. Callers that package a single HTML can pass a
        temporary ``database_path`` and embed the returned report after
        validation.
        """

        from .sqlite_report import SCHEMA_VERSION, write_sqlite_report

        if database_path is None:
            database_path = os.path.join(
                self._html_dir, f"{self.module_name}_address_map.sqlite"
            )
        if schema_version is None:
            schema_version = SCHEMA_VERSION
        return write_sqlite_report(
            self, database_path, schema_version=schema_version
        )

    def report_single_html(
        self,
        viewer_template_path=None,
        output_html_path=None,
        database_path=None,
        *,
        schema_version=None,
    ):
        """Generate SQLite and package it into one offline viewer HTML.

        When ``database_path`` is omitted the intermediate SQLite file is
        removed after successful (or failed) packaging. Passing an explicit
        path keeps it, which is useful for schema inspection and size reports.
        """

        import tempfile
        from pathlib import Path

        from .single_html_report import (
            default_viewer_template_path,
            package_single_html,
        )

        html_dir = Path(self._html_dir).expanduser().resolve()
        html_dir.mkdir(parents=True, exist_ok=True)
        if output_html_path is None:
            output_html_path = html_dir / f"{self.module_name}_address_map.html"
        if viewer_template_path is None:
            viewer_template_path = default_viewer_template_path()

        owns_database = database_path is None
        if owns_database:
            descriptor, temporary_name = tempfile.mkstemp(
                prefix=f".{self.module_name}_address_map.",
                suffix=".sqlite",
                dir=str(html_dir),
            )
            os.close(descriptor)
            os.unlink(temporary_name)
            database_path = temporary_name

        try:
            self.report_sqlite(database_path, schema_version=schema_version)
            return package_single_html(
                database_path,
                viewer_template_path,
                output_html_path,
            )
        finally:
            if owns_database:
                try:
                    Path(database_path).unlink()
                except FileNotFoundError:
                    pass
    

    # total ========================================
    def generate(
        self,
        path=None,
        gen_doc=False,
        check_ralf=False,
        viewer_template_path=None,
        report_viewer=True,
        sqlite_schema_version=None,
    ):
        if path != None:
            self.path = path
        self.report_json(gen_doc)
        if report_viewer:
            self.report_single_html(
                viewer_template_path,
                schema_version=sqlite_schema_version,
            )
        self.report_ralf()
        if check_ralf:  self.check_ralf()
        self.report_chead()
        self.report_vhead()

    def check(self, path=None):
        if path != None:
            self.path = path
        self.check_json()
        self.check_ralf()
        self.check_chead()
        self.check_vhead()
        

    #########################################
    # tablelike support
    #########################################
    def regspace(self, name,size,description='',path='./',bus_width=APG_BUS_WIDTH,software_interface='apb', offset=0):
        from .RegSpace import RegSpace

        u_ss = RegSpace(name=name, size=size, description=description, path=path, bus_width=bus_width, software_interface=software_interface)
        u_ss.offset = offset
        u_ss.father = self 
        return u_ss
    
    def addrspace(self, sub_space, offset, name):
        self.add(sub_space, offset, name)
        return self

    #########################################
    # matrix cfg
    #########################################
    def generate_matrix_excel(self, path=None):
        if path is not None:
            self.path = path
        os.makedirs(self._json_dir, exist_ok=True)

        workbook = openpyxl.Workbook()
        master_sheet = workbook.active
        master_sheet.title = "master"
        master_sheet.append(list(master_mapping.keys()))
        for values in self.report_master().values():
            master_sheet.append(values)

        slave_sheet = workbook.create_sheet(title="slave")
        slave_sheet.append(list(slave_mapping.keys()))
        for values in self.report_slave().values():
            slave_sheet.append(values)

        interconnect_sheet = workbook.create_sheet(title="interconnection")
        mapping = self.report_interconnect()
        interconnect_sheet.append(['name'] + list(mapping['name']))
        for key, values in mapping.items():
            if key != 'name':
                interconnect_sheet.append([key] + list(values))

        workbook.save(self.matrix_path)

    def report_interconnect(self):
        interconnect_set = set()
        for sub_matrix in self.matrix_list:
            for values in sub_matrix.report_interconnect().values():
                interconnect_set.update(values)

        interconnect_list = sorted(interconnect_set)
        result = {'name': interconnect_list}
        for sub_matrix in self.matrix_list:
            slaves = list(sub_matrix.report_interconnect().values())[0]
            result[sub_matrix.module_name] = [slave in slaves for slave in interconnect_list]
        return result

    def report_master(self):
        result = {}
        for sub_matrix in self.matrix_list:
            result.update(sub_matrix.report_master_matrix())
        return result

    def report_slave(self):
        result = {}
        for sub_matrix in self.matrix_list:
            result.update(sub_matrix.report_slave_matrix())
        return result

    def report_matrix(self, path=None):
        if path is not None:
            self.path = path
        os.makedirs(self._json_dir, exist_ok=True)
        matrix_json = [sub_matrix.report_json_core() for sub_matrix in self.matrix_list]
        with open(self.matrix_json_path, 'w') as output:
            json.dump(matrix_json, output, ensure_ascii=False, indent=2)

    ############################
    # check
    ############################


    def check_ralf(self):
        self.recursive_check_ralf_core()


    def recursive_check_ralf_core(self):
        for ss in self.sub_space_list:
            if hasattr(ss,'report_ralf_core'):
                ss.check_ralf()
            else:
                ss.recursive_check_ralf_core()


    def check_vhead(self):
        check_dir = os.path.join(self._vhead_dir, 'vcs_check')
        file_path = os.path.join(self._vhead_dir, 'all.vh')
        os.makedirs(check_dir, exist_ok=True)
        
        print("\n################################################################################")
        print("[Check vhead] Check vhead: %s"% file_path)
        print("################################################################################\n")
        
        command_vhead = f'vcs -full64 -sverilog -cpp g++-4.8 -cc gcc-4.8 -LDFLAGS -Wl,--no-as-needed +lint=PCWM -debug_access+all -o {check_dir}/simv -Mdir={check_dir}/csrc {file_path} | tee {check_dir}/vcs.log'
        os.system(command_vhead)
        # with open(f'{check_dir}/vcs.log','r') as f:
        #     if not re.search(r'simv\sup\sto\sdate', f.readlines()[-1]): 
        #         raise Exception("vhead check error occur, log path:%s"% os.path.abspath(f'{check_dir}/vcs.log'))
        if not os.path.exists(f'{check_dir}/simv'):
            raise Exception("vhead check error occur, log path:%s"% os.path.abspath(f'{check_dir}/vcs.log'))
        print("[Check vhead] vhead check output log: %s"% os.path.abspath(f'{check_dir}/vcs.log'))
    

    def check_json(self):
        json_path = os.path.join(self._html_dir, 'data.json')
        try:
            print("\n################################################################################")
            print("[Check Json] Check Json: %s"% json_path)
            print("################################################################################\n")
            json_file = open(json_path, 'r')
            json.load(json_file)

            print("[Check Json] json file correct")
        except:
            raise Exception("[Check Json] Error in Json file")
        
    
    def check_chead(self):
        file_path = os.path.join(self._chead_dir,'all.h')
        print("\n################################################################################")
        print("[Check chead] Check chead: %s"% file_path)
        print("################################################################################\n")
        
        if os.system('gcc -include stdint.h %s' % file_path) !=0:
            raise Exception('c head compile error.')
        print("[Check chead]  c head correct")
